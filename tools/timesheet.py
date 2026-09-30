"""Personal time sheet (Stundendokumentation) based on the course CSV template.

The file lives in stundendoku/ (gitignored – the repository is public).

Usage:
    python3 tools/timesheet.py init "Nachname Vorname"
    python3 tools/timesheet.py set <date> <hours> "<Arbeitsgegenstand>" [--comment ...] [--append]
    python3 tools/timesheet.py show [--all]
    python3 tools/timesheet.py missing [--hook]
    python3 tools/timesheet.py export          # writes an .xlsx next to the CSV

<date>: YYYY-MM-DD, DD.MM.YYYY, "heute" or "gestern".
"""

import argparse
import csv
import json
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "data" / "assets" / "Stundendokumentation_Vorlage.csv"
SHEET_DIR = ROOT / "stundendoku"
CSV_DATE = "%m/%d/%Y"  # date format used by the template
NAME_ROW, HEADER = 1, ["Datum", "Stunden", "Arbeitsgegenstand", "Kommentar"]


def sheet_path() -> Path:
    files = sorted(SHEET_DIR.glob("Stundendokumentation_*.csv"))
    if len(files) != 1:
        sys.exit(
            f"Expected exactly one Stundendokumentation_*.csv in {SHEET_DIR} (found {len(files)}). "
            'Run: python3 tools/timesheet.py init "Nachname Vorname"'
        )
    return files[0]


def read(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8", newline="") as f:
        return [row + [""] * (4 - len(row)) for row in csv.reader(f)]


def write(path: Path, rows: list[list[str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(rows)


def parse_date(text: str) -> date:
    today = date.today()
    shortcuts = {"heute": today, "today": today, "gestern": today - timedelta(days=1)}
    if text.lower() in shortcuts:
        return shortcuts[text.lower()]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", CSV_DATE):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    sys.exit(f"Unknown date '{text}' – use YYYY-MM-DD, DD.MM.YYYY, heute or gestern.")


def row_date(row: list[str]) -> date | None:
    try:
        return datetime.strptime(row[0], CSV_DATE).date()
    except ValueError:
        return None


def hours(row: list[str]) -> float:
    try:
        return float(row[1].replace(",", "."))
    except ValueError:
        return 0.0


def fmt_hours(value: float) -> str:
    return f"{round(value, 2):g}"


def sum_index(rows: list[list[str]]) -> int:
    return next(i for i, r in enumerate(rows) if r[0].strip() == "Summe")


def entries(rows: list[list[str]]) -> list[list[str]]:
    return [r for r in rows if row_date(r) and r[1].strip()]


def update_sum(rows: list[list[str]]) -> float:
    total = sum(hours(r) for r in rows if row_date(r))
    rows[sum_index(rows)][1] = fmt_hours(total)
    return total


def cmd_init(args: argparse.Namespace) -> None:
    if not TEMPLATE.exists():
        sys.exit(f"Template missing: {TEMPLATE} (download it from ILIAS).")
    existing = list(SHEET_DIR.glob("Stundendokumentation_*.csv"))
    if existing:
        sys.exit(f"Already initialised: {existing[0].relative_to(ROOT)}")
    rows = read(TEMPLATE)
    rows[NAME_ROW][0] = args.name
    SHEET_DIR.mkdir(exist_ok=True)
    path = SHEET_DIR / f"Stundendokumentation_{args.name.replace(' ', '_')}.csv"
    write(path, rows)
    print(f"Created {path.relative_to(ROOT)}")


def cmd_set(args: argparse.Namespace) -> None:
    path = sheet_path()
    rows = read(path)
    day = parse_date(args.date)
    key = day.strftime(CSV_DATE)
    idx = next((i for i, r in enumerate(rows) if r[0] == key), None)
    if idx is None:
        later = [i for i, r in enumerate(rows) if row_date(r) and row_date(r) > day]
        idx = later[0] if later else sum_index(rows)
        rows.insert(idx, [key, "", "", ""])

    row = rows[idx]
    new_hours, text = args.hours, args.text.strip()
    if row[1].strip() and not args.append:
        sys.exit(f"{day:%d.%m.%Y} already has an entry: {row[1]} h – {row[2]}. Use --append.")
    if args.append and row[1].strip():
        new_hours += hours(row)
        text = f"{row[2].strip()}; {text}" if row[2].strip() else text
    row[1], row[2] = fmt_hours(new_hours), text
    if args.comment is not None:
        row[3] = args.comment
    total = update_sum(rows)
    write(path, rows)
    print(f"{day:%d.%m.%Y}: {row[1]} h – {row[2]}\nSumme: {fmt_hours(total)} h")


def cmd_show(args: argparse.Namespace) -> None:
    rows = read(sheet_path())
    shown = entries(rows)
    if not args.all:
        shown = shown[-7:]
    for r in shown:
        print(f"{row_date(r):%a %d.%m.}  {r[1]:>5} h  {r[2]}")
    print(f"Summe: {rows[sum_index(rows)][1] or 0} h")


def missing_days() -> list[date]:
    """Past days with own commits (any branch) but no time sheet entry."""
    email = subprocess.run(
        ["git", "config", "user.email"], capture_output=True, text=True, cwd=ROOT
    ).stdout.strip()
    log = subprocess.run(
        ["git", "log", "--all", f"--author={email}", "--format=%ad", "--date=short"],
        capture_output=True,
        text=True,
        cwd=ROOT,
    ).stdout.split()
    commit_days = {date.fromisoformat(d) for d in log}
    done = {row_date(r) for r in entries(read(sheet_path()))}
    return sorted(d for d in commit_days - done if d < date.today())


def cmd_missing(args: argparse.Namespace) -> None:
    if args.hook:
        try:
            if not list(SHEET_DIR.glob("Stundendokumentation_*.csv")):
                message = "Stundendoku noch nicht eingerichtet – `/stunden init` ausführen."
                if not TEMPLATE.exists():
                    return
            else:
                days = missing_days()
                if not days:
                    return
                listed = ", ".join(f"{d:%d.%m.}" for d in days[-5:])
                message = f"Stundendoku: Einträge fehlen für {listed} – mit `/stunden` nachtragen."
            print(json.dumps({
                "systemMessage": message,
                "hookSpecificOutput": {"hookEventName": "SessionStart",
                                       "additionalContext": message},
            }))  # fmt: skip
        except Exception:  # noqa: BLE001 – a reminder must never break a session
            pass
        return
    days = missing_days()
    print("\n".join(f"{d:%Y-%m-%d}" for d in days) or "Nichts offen.")


def cmd_export(args: argparse.Namespace) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        sys.exit("openpyxl missing – run via `make stunden` or `uv run`.")
    path = sheet_path()
    rows = read(path)
    update_sum(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Stunden"
    for r in rows:
        day = row_date(r)
        value = hours(r) if (day or r[0] == "Summe") and r[1].strip() else r[1]
        ws.append([day or r[0], value, r[2], r[3]])
        if day:
            ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"
        if r[0] in ("Summe", HEADER[0]) or ws.max_row == 1:
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
    for col, width in zip("ABCD", (12, 9, 90, 40), strict=True):
        ws.column_dimensions[col].width = width
    out = path.with_suffix(".xlsx")
    wb.save(out)
    print(f"Exported {out.relative_to(ROOT)} (Summe: {rows[sum_index(rows)][1]} h)")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("init", help="create personal sheet from the template")
    p.add_argument("name", help='"Nachname Vorname"')
    p = sub.add_parser("set", help="write the entry for one day")
    p.add_argument("date")
    p.add_argument("hours", type=lambda s: float(s.replace(",", ".")))
    p.add_argument("text")
    p.add_argument("--comment")
    p.add_argument("--append", action="store_true", help="add to an existing entry")
    p = sub.add_parser("show", help="show recent entries and the sum")
    p.add_argument("--all", action="store_true")
    p = sub.add_parser("missing", help="days with commits but without entry")
    p.add_argument("--hook", action="store_true", help=argparse.SUPPRESS)
    sub.add_parser("export", help="export to .xlsx")
    args = parser.parse_args()
    {"init": cmd_init, "set": cmd_set, "show": cmd_show, "missing": cmd_missing,
     "export": cmd_export}[args.command](args)  # fmt: skip


if __name__ == "__main__":
    main()
